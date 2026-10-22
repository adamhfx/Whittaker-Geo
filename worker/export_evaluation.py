"""Build an export from the same validated values used by the live dashboard."""
from copy import copy
from pathlib import Path
from typing import Any

import openpyxl

from .model_preview import PreviewError, _hidden_columns, _views, export_preview_results


def export_evaluation(template: Path, manifest: dict[str, Any], result: dict[str, Any], target: Path) -> None:
    # Read only source geometry. Never copy the original workbook package,
    # formulas, private sheets, cached values, or drawing relationships.
    layout = openpyxl.load_workbook(template, data_only=False)
    clean = openpyxl.Workbook()
    clean.remove(clean.active)
    try:
        views = _views(manifest)
        if len(result['outputSheets']) != len(views):
            raise PreviewError('RESULT_INVALID')
        for view, values in zip(views, result['outputSheets']):
            source = layout[view['sheet']]
            destination = clean.create_sheet(view['sheet'])
            if source.sheet_state != 'visible' or values['name'] != view['sheet']:
                raise PreviewError('OUTPUT_MAPPING_INVALID')
            for name, dimension in source.column_dimensions.items():
                destination.column_dimensions[name] = copy(dimension)
            left, top, right, bottom = view['bounds']
            hidden = _hidden_columns(source)
            rows = [r for r in range(top, bottom + 1) if not source.row_dimensions[r].hidden]
            columns = [c for c in range(left, right + 1) if c not in hidden]
            if len(values['rows']) != len(rows):
                raise PreviewError('RESULT_INVALID')
            for r in range(top, bottom + 1):
                destination.row_dimensions[r] = copy(source.row_dimensions[r])
            for row_index, row in zip(rows, values['rows']):
                actual_columns = [cell['column'] for cell in row['cells']]
                # The live contract omits blank cells but preserves every
                # visible row. Accept a sparse, ordered subset only.
                if actual_columns != sorted(set(actual_columns)) or not set(actual_columns).issubset(columns):
                    raise PreviewError('RESULT_INVALID')
                for item in row['cells']:
                    cell = destination.cell(row_index, item['column'], item['value'])
                    if isinstance(item['value'], str):
                        cell.data_type = 's'
                    cell.number_format = item.get('format', 'General')
        presentation = {key: result[key] for key in ('summary', 'charts') if key in result}
        export_preview_results(clean, manifest, target, presentation=presentation)
    finally:
        clean.close()
        layout.close()
