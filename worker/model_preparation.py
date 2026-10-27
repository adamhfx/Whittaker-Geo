"""Private candidate preparation from an existing calculator contract.

No uploaded instructions are executed. The source is immutable; deterministic
runtime copies and explicit mapping changes are reviewed before registration.
"""
from __future__ import annotations

import copy
import hashlib
import json
import math
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries, range_to_tuple

from .engine import ModelError, inject_inputs, validate_inputs
from .formula_compatibility import inline_nested_let, FormulaCompatibilityError
from .model_preview import _hidden_columns, _views, _presentation_specs

NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
REL = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
CELL = re.compile(r'^[A-Z]{1,3}[1-9][0-9]{0,6}$')
FORBIDDEN = re.compile(r'(?i)(?:^|[^A-Z0-9_.])(?:_xlfn\.)?(?:DDE|WEBSERVICE|RTD|CALL|REGISTER\.ID|SQL\.REQUEST|STOCKHISTORY|IMAGE)\s*\(')


class PreparationError(ValueError):
    pass


def check_package(path: Path) -> None:
    if path.stat().st_size > 20 * 1024 * 1024:
        raise PreparationError('The workbook exceeds 20 MB.')
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        total_cells = 0
        if len(entries) > 2000 or sum(item.file_size for item in entries) > 100 * 1024 * 1024:
            raise PreparationError('The expanded workbook exceeds the supported size.')
        for item in entries:
            name = item.filename.lower()
            if (name.startswith('/') or '..' in name.split('/') or '\\' in name or item.file_size > 25 * 1024 * 1024
                or item.compress_size and item.file_size / item.compress_size > 100
                or any(part in name for part in ('vbaproject', '/embeddings/', 'externallinks/'))
                or re.search(r'\.(exe|dll|js|vbs|bat|cmd|ps1)$', name)):
                raise PreparationError('The workbook contains unsupported external or executable content.')
            if name.endswith(('.xml', '.rels')):
                data = archive.read(item)
                if re.search(br'<!DOCTYPE|<!ENTITY', data, re.I):
                    raise PreparationError('Workbook XML declarations are unsupported.')
                document = ET.fromstring(data)
                if name.endswith('.rels') and any(rel.attrib.get('TargetMode', '').lower() == 'external' for rel in document):
                    raise PreparationError('External workbook connections are unsupported.')
                if name == 'xl/connections.xml' and list(document):
                    raise PreparationError('Workbook data connections are unsupported.')
                if name.startswith('xl/worksheets/'):
                    cells = document.findall('.//m:sheetData/m:row/m:c', NS)
                    total_cells += len(cells)
                    if total_cells > 300_000:
                        raise PreparationError('The workbook exceeds the supported cell count.')
                    for cell in cells:
                        formula = cell.find('m:f', NS)
                        if formula is not None:
                            value = formula.text or ''
                            unquoted = re.sub(r'"(?:""|[^"])*"|\'(?:\'\'|[^\'])*\'', '', value)
                            if FORBIDDEN.search(value) or '|' in unquoted:
                                raise PreparationError('An external-data or executable formula is unsupported.')


def _color(color) -> tuple:
    return (color.type, str(color.value), color.tint) if color is not None else ()


def _input_style(cell) -> tuple:
    return (cell.fill.patternType, _color(cell.fill.fgColor), _color(cell.font.color), cell.protection.locked)


def _scalar(value):
    if value is None or isinstance(value, str) and len(value) <= 200:
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value):
        return value
    raise PreparationError('An input default is not a supported scalar value.')


def _anchor(sheet, cell, inputs: set[str]) -> tuple[str, ...]:
    """Compare nearby fixed labels/years, excluding other editable values."""
    labels = []
    positions = [(cell.row, column) for column in range(max(1, cell.column - 8), cell.column)]
    positions += [(row, cell.column) for row in range(max(1, cell.row - 4), cell.row)]
    for row, column in positions:
        neighbour = sheet.cell(row, column)
        if neighbour.coordinate in inputs or neighbour.data_type == 'f' or neighbour.value is None:
            continue
        if isinstance(neighbour.value, (str, int, float)):
            labels.append(re.sub(r'\s+', ' ', str(neighbour.value)).strip().casefold())
    return tuple(labels)


def _text(chart_title, default: str) -> str:
    try:
        paragraphs = chart_title.tx.rich.p
        value = ''.join(run.t for paragraph in paragraphs for run in paragraph.r)
        return value.strip()[:120] or default
    except (AttributeError, TypeError):
        return default


def _chart_specs(book, cached, views: list[dict]) -> list[dict]:
    charts = []
    for view in views:
        for chart in book[view['sheet']]._charts:
            kind_name = type(chart).__name__
            if kind_name not in {'LineChart', 'BarChart'}:
                raise PreparationError('A published chart type needs a supported line or column chart.')
            series = []
            for index, item in enumerate(chart.series):
                category = getattr(item, 'cat', None)
                if category is not None and any(getattr(category, name, None) is not None for name in ('numRef','strRef','numLit','strLit','multiLvlStrRef')):
                    raise PreparationError('A published chart uses explicit category labels that need developer mapping. Automatic preparation currently supports sequential point labels.')
                reference = getattr(getattr(item, 'val', None), 'numRef', None)
                if not reference or not reference.f or '[' in reference.f:
                    raise PreparationError('A published chart needs workbook-local numeric series.')
                sheet, bounds = range_to_tuple(reference.f)
                if sheet not in book.sheetnames:
                    raise PreparationError('A chart series worksheet is missing.')
                left, top, right, bottom = bounds
                if (right != left and bottom != top) or max(right-left+1, bottom-top+1) > 100:
                    raise PreparationError('A chart series exceeds the supported single-row/column range.')
                label = getattr(getattr(item, 'tx', None), 'v', None)
                label_ref = getattr(getattr(item, 'tx', None), 'strRef', None)
                if label_ref and label_ref.f:
                    label_sheet, label_bounds = range_to_tuple(label_ref.f)
                    a, b, c, d = label_bounds
                    if label_sheet not in book.sheetnames or a != c or b != d:
                        raise PreparationError('A chart series label is not a single workbook cell.')
                    label = cached[label_sheet].cell(b, a).value or book[label_sheet].cell(b, a).value
                if not isinstance(label, str) or label.startswith('='):
                    label = f'Series {index + 1}'
                series.append({'name': label[:100], 'sheet': sheet,
                               'range': f'{openpyxl.utils.get_column_letter(left)}{top}:{openpyxl.utils.get_column_letter(right)}{bottom}'})
            charts.append({'title': _text(chart.title, f'Chart {len(charts)+1}'),
                           'kind': 'line' if kind_name == 'LineChart' else 'column', 'series': series})
    if len(charts) > 8:
        raise PreparationError('The workbook has more than eight published charts.')
    return charts


def inspect_candidate(source: Path, base_template: Path, base: dict, version: str, mapping: dict | None = None) -> tuple[dict, dict]:
    check_package(source)
    mapping = mapping or {}
    if not isinstance(mapping, dict):
        raise PreparationError('The input mapping is invalid.')
    book = openpyxl.load_workbook(source, data_only=False, keep_links=False)
    cached = openpyxl.load_workbook(source, data_only=True, keep_links=False)
    previous = openpyxl.load_workbook(base_template, data_only=False, keep_links=False)
    try:
        sheet_name = base.get('inputSheet', 'Inputs')
        if sheet_name not in book.sheetnames or book[sheet_name].sheet_state != 'visible':
            raise PreparationError('The configured input worksheet must be present and visible.')
        current_sheet, old_sheet = book[sheet_name], previous[sheet_name]
        fields = base['fields']
        old_cells = {field['cell'].upper() for field in fields}
        styles = {_input_style(old_sheet[cell]) for cell in old_cells}
        candidates = {cell.coordinate for cell in current_sheet._cells.values()
                      if cell.data_type != 'f' and not cell.protection.locked and _input_style(cell) in styles}
        old_candidates = {cell.coordinate for cell in old_sheet._cells.values()
                          if cell.data_type != 'f' and not cell.protection.locked and _input_style(cell) in styles}
        candidate_anchors: dict[tuple, list[str]] = {}
        for address in candidates:
            anchor = _anchor(current_sheet, current_sheet[address], candidates)
            if anchor: candidate_anchors.setdefault(anchor, []).append(address)
        allowed_keys = {'fieldCells', 'fieldMetadata', 'addedFields', 'ignoredCells', 'outputRanges', 'summaryCells', 'feedbackCells', 'addedFeedback', 'ignoredFeedback', 'help', 'guideReviewed', 'testCases'}
        if set(mapping) - allowed_keys:
            raise PreparationError('The mapping contains unsupported settings.')
        overrides = mapping.get('fieldCells', {})
        if not isinstance(overrides, dict) or set(overrides) - {field['id'] for field in fields}:
            raise PreparationError('The mapping refers to an unknown input field.')
        metadata = mapping.get('fieldMetadata', {})
        if not isinstance(metadata, dict) or set(metadata) - {field['id'] for field in fields}:
            raise PreparationError('Input metadata refers to an unknown field.')
        operational = ('id', 'name', 'description', 'inputSheet', 'groups', 'fields', 'outputs', 'multipliers', 'inputRules', 'presentation', 'inputFeedback')
        manifest = {key: copy.deepcopy(base[key]) for key in operational if key in base}
        manifest['version'], manifest['releaseStatus'] = version, 'review_required'
        id_map = {field['id']: field['id'].replace(f"{base['id']}@{base['version']}:", f"{base['id']}@{version}:", 1) for field in fields}
        report: dict[str, Any] = {'fields': [], 'unmappedCells': [], 'issues': [], 'changes': [], 'guideChanges': []}
        used: set[str] = set()
        retained = []
        for field in manifest['fields']:
            original_id = field['id']
            old_address = field['cell'].upper()
            if original_id in overrides and overrides[original_id] is None:
                report['changes'].append({'label': field['label'], 'kind': 'removed_input'})
                continue
            explicit = original_id in overrides
            old_anchor = _anchor(old_sheet, old_sheet[old_address], old_candidates)
            address = overrides.get(original_id, old_address)
            if not explicit and (address not in candidates or _anchor(current_sheet, current_sheet[address], candidates) != old_anchor):
                matches = candidate_anchors.get(old_anchor, [])
                address = matches[0] if len(matches) == 1 else address
                if len(matches) != 1:
                    report['issues'].append(f"Review the worksheet cell for {field['label']}.")
            if not isinstance(address, str) or not CELL.fullmatch(address.upper()) or address.upper() not in candidates or address.upper() in used:
                report['issues'].append(f"Choose a unique, unlocked non-formula input cell for {field['label']}.")
                address = old_address
            address = address.upper()
            used.add(address)
            field['id'], field['cell'] = id_map[original_id], address
            if field.get('dependsOn') in id_map: field['dependsOn'] = id_map[field['dependsOn']]
            if original_id in metadata:
                edits = metadata[original_id]
                if not isinstance(edits, dict) or set(edits) - {'label','unit','min','max','exclusiveMin','exclusiveMax','integer','required','blankOnNew','year'}:
                    raise PreparationError('Input metadata contains unsupported properties.')
                for key, setting in edits.items():
                    if key in {'label','unit'} and (not isinstance(setting, str) or len(setting)>200):
                        raise PreparationError('Input labels and units must be short text.')
                    elif key in {'integer','required','blankOnNew'} and not isinstance(setting, bool):
                        raise PreparationError('Input validation selections must be boolean.')
                    elif key in {'min','max','exclusiveMin','exclusiveMax','year'}:
                        if setting is None:
                            field.pop(key, None);continue
                        if isinstance(setting,bool) or not isinstance(setting,(int,float)) or not math.isfinite(setting):
                            raise PreparationError('Input bounds must be finite numbers.')
                    field[key]=setting
            value = None if current_sheet[address].data_type == 'f' else _scalar(current_sheet[address].value)
            field['default'] = value
            if value != next(item['default'] for item in fields if item['id'] == original_id):
                report['changes'].append({'label': field['label'], 'kind': 'default', 'before': next(item['default'] for item in fields if item['id'] == original_id), 'after': value})
            if old_address != address: report['changes'].append({'label': field['label'], 'kind': 'moved_input', 'before': old_address, 'after': address})
            report['fields'].append({'id': original_id, 'label': field['label'], 'group': field['group'], 'type': field['type'], 'unit': field['unit'], 'baseCell': old_address, 'cell': address, 'default': value,
                **{key:field[key] for key in ('min','max','exclusiveMin','exclusiveMax','integer','required','blankOnNew','year') if key in field}})
            retained.append(field)
        manifest['fields'] = retained
        added = mapping.get('addedFields', [])
        if not isinstance(added, list) or len(added) > 100:
            raise PreparationError('At most 100 new inputs may be reviewed in one revision.')
        for item in added:
            if not isinstance(item, dict): raise PreparationError('A new input definition is invalid.')
            address = str(item.get('cell', '')).upper()
            label, unit, kind = item.get('label'), item.get('unit', ''), item.get('type')
            if (address not in candidates or address in used or not isinstance(label, str) or not 1 <= len(label) <= 200
                or not isinstance(unit, str) or len(unit) > 100 or kind not in {'number', 'text', 'select'}):
                raise PreparationError('A new input needs a unique editable cell, label, unit and supported type.')
            field = {'id': f"{base['id']}@{version}:new:{address.lower()}", 'cell': address, 'label': label, 'unit': unit,
                     'type': kind, 'group': 'additional-inputs', 'required': True, 'blankOnNew': True, 'default': _scalar(current_sheet[address].value)}
            if kind == 'select':
                options = item.get('options')
                if not isinstance(options, list) or not 1 <= len(options) <= 100 or any(not isinstance(option, str) or len(option) > 200 for option in options):
                    raise PreparationError('A new choice input needs a bounded list of exact workbook option values.')
                field['options'] = options
            manifest['fields'].append(field)
            used.add(address)
            report['changes'].append({'label': label, 'kind': 'added_input', 'after': address})
        if added and not any(group['id']=='additional-inputs' for group in manifest['groups']):
            manifest['groups'].append({'id': 'additional-inputs', 'name': 'Additional inputs'})
        if len(manifest['fields']) > 500:
            raise PreparationError('The calculator exceeds 500 editable inputs.')
        ignored = mapping.get('ignoredCells', [])
        if not isinstance(ignored, list) or any(not isinstance(cell, str) for cell in ignored):
            raise PreparationError('Ignored cell selections are invalid.')
        ignored = ignored + [field['cell'].upper() for field in fields if field['id'] in overrides and overrides[field['id']] is None]
        new_unmapped = candidates - used - (old_candidates - old_cells) - set(ignored)
        report['unmappedCells'] = sorted(new_unmapped)
        if new_unmapped: report['issues'].append('Define or explicitly exclude the new editable workbook cells.')
        remaining_ids = {field['id'] for field in manifest['fields']}
        for field in manifest['fields']:
            if field.get('dependsOn') and field['dependsOn'] not in remaining_ids:
                report['issues'].append(f"The dependency for {field['label']} was removed and needs model setup.")
        manifest['multipliers'] = [id_map[key] for key in base.get('multipliers', []) if id_map.get(key) in remaining_ids]
        for rule in manifest.get('inputRules', {}).get('ascending', []):
            rule['fields'] = [id_map.get(key, key) for key in rule['fields']]
            if any(key not in remaining_ids for key in rule['fields']):
                report['issues'].append('A removed input is still required by a cross-field validation rule.')
        output_ranges = mapping.get('outputRanges', {})
        if not isinstance(output_ranges, dict) or set(output_ranges)-{view['id'] for view in manifest['outputs']['views']}:
            raise PreparationError('Output ranges are invalid.')
        for view in manifest['outputs']['views']:
            if view['id'] in output_ranges:
                view['sourceRange'] = output_ranges[view['id']]
            if view['sheet'] not in book.sheetnames or book[view['sheet']].sheet_state != 'visible':
                raise PreparationError('A configured OUTPUT worksheet is missing or hidden.')
        views = _views(manifest)
        if manifest.get('presentation') is not None:
            summary_overrides = mapping.get('summaryCells', {})
            if not isinstance(summary_overrides, dict): raise PreparationError('Summary mappings are invalid.')
            report['summaryCells'] = []
            for index, item in enumerate(manifest['presentation']['summary']):
                original_sheet, original_cell = item['sheet'], item['valueCell']
                explicit = str(index) in summary_overrides
                if explicit:
                    edit = summary_overrides[str(index)]
                    if not isinstance(edit, dict) or set(edit)-{'sheet','cell','label','unit'}:
                        raise PreparationError('A summary mapping is invalid.')
                    item['sheet'], item['valueCell'] = edit.get('sheet',original_sheet), edit.get('cell',original_cell)
                    for key in ('label','unit'):
                        if key in edit:
                            if not isinstance(edit[key],str) or len(edit[key])>(120 if key=='label' else 100):raise PreparationError('A summary label or unit is invalid.')
                            item[key]=edit[key]
                if item['sheet'] not in {view['sheet'] for view in views} or not isinstance(item['valueCell'],str) or not CELL.fullmatch(item['valueCell'].upper()):
                    raise PreparationError('Summary values must use a configured visible OUTPUT worksheet.')
                item['valueCell']=item['valueCell'].upper()
                old_anchor=_anchor(previous[original_sheet],previous[original_sheet][original_cell],set())
                new_anchor=_anchor(book[item['sheet']],book[item['sheet']][item['valueCell']],set())
                if not explicit and old_anchor!=new_anchor:
                    report['issues'].append(f"Review the summary value cell for {item['label']}.")
                report['summaryCells'].append({'index':index,'label':item['label'],'sheet':item['sheet'],'cell':item['valueCell'],'unit':item.get('unit') or ''})
            manifest['presentation']['charts'] = _chart_specs(book, cached, views)
            _presentation_specs(manifest)
        feedback_overrides=mapping.get('feedbackCells',{})
        if not isinstance(feedback_overrides,dict) or set(feedback_overrides)-{item['id'] for item in manifest.get('inputFeedback',[])}:
            raise PreparationError('Interim feedback mappings are invalid.')
        report['feedbackCells']=[]
        for item in manifest.get('inputFeedback', []):
            original_cell=item['cell']
            item['cell'] = feedback_overrides.get(item['id'], item['cell'])
            if not isinstance(item['cell'], str) or not CELL.fullmatch(item['cell'].upper()):
                raise PreparationError('An interim feedback cell is invalid.')
            item['cell']=item['cell'].upper()
            if item['id'] not in feedback_overrides and _anchor(old_sheet,old_sheet[original_cell],old_candidates)!=_anchor(current_sheet,current_sheet[item['cell']],candidates):
                report['issues'].append(f"Review the calculated feedback cell for {item['label']}.")
            report['feedbackCells'].append({'id':item['id'],'label':item['label'],'group':item['group'],'unit':item.get('unit',''),'cell':item['cell']})
        added_feedback=mapping.get('addedFeedback',[])
        if not isinstance(added_feedback,list) or len(added_feedback)>100:raise PreparationError('Additional feedback definitions are invalid.')
        for item in added_feedback:
            if not isinstance(item,dict) or not isinstance(item.get('cell'),str) or not CELL.fullmatch(item['cell'].upper()) or not isinstance(item.get('label'),str) or not 1<=len(item['label'])<=160 or item.get('group') not in {group['id'] for group in manifest['groups']} or not isinstance(item.get('unit',''),str):
                raise PreparationError('Each additional feedback value needs a cell, label, group and unit.')
            address=item['cell'].upper()
            if address in used:raise PreparationError('A feedback value cannot overlap an editable input.')
            manifest.setdefault('inputFeedback',[]).append({'id':'feedback-new-'+address.lower(),'cell':address,'label':item['label'],'group':item['group'],'unit':item.get('unit','')[:100]})
        if 'User Guide' in book.sheetnames and 'User Guide' in previous.sheetnames:
            old_guide = {cell.coordinate: cell.value for cell in previous['User Guide']._cells.values() if isinstance(cell.value, str) and cell.data_type != 'f'}
            new_guide = {cell.coordinate: cell.value for cell in book['User Guide']._cells.values() if isinstance(cell.value, str) and cell.data_type != 'f'}
            report['guideChanges'] = [{'cell': key, 'before': str(old_guide.get(key, ''))[:2000], 'after': str(new_guide.get(key, ''))[:2000]} for key in sorted(set(old_guide) | set(new_guide)) if old_guide.get(key) != new_guide.get(key)][:100]
            if report['guideChanges'] and mapping.get('guideReviewed') is not True:
                report['issues'].append('Review the changed User Guide and the published input notes.')
        help_overrides = mapping.get('help', {})
        if not isinstance(help_overrides, dict): raise PreparationError('Guide edits are invalid.')
        for item in manifest['fields'] + manifest['groups']:
            source_id = next((key for key, value in id_map.items() if value == item['id']), item['id'])
            if source_id in help_overrides:
                note = help_overrides[source_id]
                if note is not None and (not isinstance(note, dict) or not isinstance(note.get('title'), str) or len(note['title']) > 200 or not isinstance(note.get('body'), str) or len(note['body']) > 12000):
                    raise PreparationError('A guide note exceeds its supported format.')
                item['help'] = note
        if not report['issues']:
            try: validate_inputs({field['id']: field['default'] for field in manifest['fields']}, manifest)
            except ModelError: report['issues'].append('The source default scenario is incomplete or invalid. Recalculate and save a complete example in Excel.')
        report['issues'] = list(dict.fromkeys(report['issues']))
        report['inputCount'] = len(manifest['fields'])
        report['candidateCells'] = [{'cell': address, 'value': _scalar(current_sheet[address].value)} for address in sorted(candidates)]
        report['outputViews'] = [{'id': view['id'], 'sheet': view['sheet'], 'range': next(item['sourceRange'] for item in manifest['outputs']['views'] if item['id'] == view['id'])} for view in views]
        report['chartCount'] = len(manifest.get('presentation', {}).get('charts', []))
        manifest['sourceSha256'] = hashlib.sha256(source.read_bytes()).hexdigest()
        return manifest, report
    finally:
        book.close(); cached.close(); previous.close()


def prepare_runtime(source: Path, target: Path, input_sheet: str) -> dict:
    injected = target.with_name('cache-cleared.xlsx')
    inject_inputs(source, injected, {}, input_sheet=input_sheet)
    count = 0
    with zipfile.ZipFile(injected) as archive, zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as output:
        for item in archive.infolist():
            data = archive.read(item)
            if item.filename.startswith('xl/worksheets/') and item.filename.endswith('.xml'):
                document = ET.fromstring(data)
                changed = False
                for cell in document.findall('.//m:sheetData/m:row/m:c', NS):
                    formula = cell.find('m:f', NS)
                    if formula is not None and formula.text:
                        try: updated, rewrites = inline_nested_let(formula.text)
                        except FormulaCompatibilityError as error: raise PreparationError(str(error)) from None
                        if rewrites:
                            formula.text = updated
                            count += rewrites
                            changed = True
                if changed: data = ET.tostring(document, encoding='utf-8', xml_declaration=True)
            metadata = zipfile.ZipInfo(item.filename, date_time=(1980, 1, 1, 0, 0, 0))
            metadata.compress_type = zipfile.ZIP_DEFLATED
            output.writestr(metadata, data, compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)
    injected.unlink()
    return {'runtimeSha256': hashlib.sha256(target.read_bytes()).hexdigest(), 'compatibilityRewrites': count}
