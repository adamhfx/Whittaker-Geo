"""Exercise the configured LibreOffice and UNO bridge with a synthetic workbook.

This is a build-time compatibility gate. It intentionally contains no client
workbook or client formula text; it verifies that the actual office process can
load, recalculate, and reset a small LET-compatible arithmetic model through
``WarmOffice``.
"""
from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
import subprocess
import tempfile
from typing import Any

import openpyxl
from openpyxl.workbook.properties import CalcProperties

from .warm_office import WarmOffice


def _make_workbook(path: Path) -> None:
    workbook = openpyxl.Workbook()
    inputs = workbook.active
    inputs.title = "UI"
    inputs["B1"] = 0
    inputs["B2"] = 0

    calculation = workbook.create_sheet("Model Calc")
    calculation.sheet_state = "hidden"
    # Excel stores newer functions with the _xlfn prefix in OOXML. Keep a LET
    # binding around a nested arithmetic expression so Calc exercises that
    # import path without depending on the candidate's transformed formulas.
    calculation["A1"] = "=_xlfn.LET(x,UI!B1,x*(5+5)+UI!B2)"

    output = workbook.create_sheet("Smoke OUTPUT")
    output["A1"] = "LET total"
    output["B1"] = "='Model Calc'!A1"
    workbook.calculation = CalcProperties(
        calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True
    )
    workbook.save(path)
    workbook.close()


def _manifest(path: Path) -> dict[str, Any]:
    return {
        "id": "runtime-smoke",
        "version": "1.0.0",
        "runtimeSha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "inputSheet": "UI",
        "fields": [
            {
                "id": "unitPrice", "cell": "UI!B1", "label": "Unit price",
                "type": "number", "required": True, "min": 0,
            },
            {
                "id": "adjustment", "cell": "UI!B2", "label": "Adjustment",
                "type": "number", "required": False, "min": 0,
            },
        ],
        "outputs": {"views": [{
            "id": "smoke-output", "label": "Runtime smoke output",
            "sheet": "Smoke OUTPUT", "sourceRange": "A1:B1",
        }]},
        "presentation": {"summary": [{
            "label": "LET total", "sheet": "Smoke OUTPUT", "valueCell": "B1",
        }]},
    }


def _read_total(result: dict[str, Any]) -> float:
    if len(result.get("outputSheets", [])) != 1:
        raise RuntimeError("Runtime smoke returned an unexpected output-sheet count.")
    output = result["outputSheets"][0]
    if output.get("name") != "Smoke OUTPUT" or len(output.get("rows", [])) != 1:
        raise RuntimeError("Runtime smoke returned an unexpected output shape.")
    cells = {cell["column"]: cell["value"] for cell in output["rows"][0]["cells"]}
    value = cells.get(2)
    summary = result.get("summary")
    if (
        not isinstance(value, (int, float)) or isinstance(value, bool)
        or not isinstance(summary, list) or len(summary) != 1
        or summary[0].get("value") != value
    ):
        raise RuntimeError("Runtime smoke did not return the calculated total.")
    return float(value)


def _probe_synthetic_cli_load(soffice_bin: str, workbook_path: Path, workdir: Path) -> None:
    """Load and convert the synthetic XLSX without the external Python-UNO bridge."""
    output_dir = workdir / "cli-output"
    output_dir.mkdir()
    profile_dir = workdir / "cli-profile"
    expected_output = output_dir / f"{workbook_path.stem}.ods"
    try:
        completed = subprocess.run(
            [
                soffice_bin, "--headless", "--nologo", "--nodefault", "--norestore",
                f"-env:UserInstallation={profile_dir.as_uri()}",
                "--convert-to", "ods", "--outdir", str(output_dir), str(workbook_path),
            ],
            capture_output=True, text=True, timeout=45, check=False,
        )
        stdout = completed.stdout
        stderr = completed.stderr
        # The fixture and all paths are temporary. Redact their root and cap
        # diagnostic output so only safe, bounded process evidence reaches CI.
        for path in {str(workdir.resolve()), workdir.resolve().as_posix()}:
            stdout = stdout.replace(path, "<synthetic-temp>")
            stderr = stderr.replace(path, "<synthetic-temp>")
        report = {
            "event": "calculation_engine_cli_load_probe",
            "exitCode": completed.returncode,
            "outputCreated": expected_output.is_file(),
            "stdout": stdout[-3000:],
            "stderr": stderr[-3000:],
        }
        print(json.dumps(report, separators=(",", ":")), flush=True)
        if completed.returncode != 0 or not expected_output.is_file():
            raise RuntimeError("Synthetic XLSX CLI load probe failed.")
    except subprocess.TimeoutExpired as error:
        def safe_text(value: str | bytes | None) -> str:
            if isinstance(value, bytes):
                value = value.decode("utf-8", errors="replace")
            if not isinstance(value, str):
                return ""
            value = value.replace(str(workdir.resolve()), "<synthetic-temp>")
            value = value.replace(workdir.resolve().as_posix(), "<synthetic-temp>")
            return value[-3000:]

        print(json.dumps({
            "event": "calculation_engine_cli_load_probe",
            "timedOut": True,
            "stdout": safe_text(error.stdout),
            "stderr": safe_text(error.stderr),
        }, separators=(",", ":")), flush=True)
        raise RuntimeError("Synthetic XLSX CLI load probe timed out.") from None


def run() -> dict[str, Any]:
    if hasattr(os, "geteuid") and os.geteuid() == 0:
        raise RuntimeError("The calculation-engine runtime smoke must run as nonroot.")
    python_bin = os.environ.get("UNO_PYTHON_BIN", "/usr/bin/python3")
    soffice_bin = os.environ.get("SOFFICE_BIN", "soffice")
    with tempfile.TemporaryDirectory(prefix="geoflow-runtime-smoke-") as temp:
        workbook_path = Path(temp) / "synthetic-model.xlsx"
        _make_workbook(workbook_path)
        manifest = _manifest(workbook_path)
        _probe_synthetic_cli_load(soffice_bin, workbook_path, Path(temp))
        office = WarmOffice(python_bin, soffice_bin, diagnostic_synthetic=True)
        try:
            cases = [
                ({"unitPrice": 2, "adjustment": 1}, 21),
                ({"unitPrice": 3, "adjustment": 7}, 37),
                # Omitting an optional value must overwrite the prior request's
                # value with its validated zero default.
                ({"unitPrice": 2}, 20),
                ({"unitPrice": 2, "adjustment": 1}, 21),
            ]
            observed = []
            for inputs, expected in cases:
                result = office.evaluate(
                    workbook_path, manifest, inputs, timeout=45, workdir=temp
                )
                value = _read_total(result)
                if value != expected:
                    raise RuntimeError(
                        f"Runtime smoke calculated {value}; expected {expected}."
                    )
                observed.append(value)
        finally:
            office.close()
    return {
        "event": "calculation_engine_runtime_smoke_passed",
        "observedTotals": observed,
        "inputResetVerified": observed == [21.0, 37.0, 20.0, 21.0],
    }


if __name__ == "__main__":
    print(json.dumps(run(), separators=(",", ":")), flush=True)
