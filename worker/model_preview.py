"""Transient, values-only model preview with a private local review CLI.

The public preview service loads only approved immutable model versions from
Supabase. The review subcommand is local-only and can inspect unapproved
candidates without opening an HTTP endpoint or writing a run.
"""
from __future__ import annotations

import argparse
from datetime import date, datetime, time as datetime_time
import hashlib
import json
import math
import re
import tempfile
import time
from pathlib import Path
from urllib.parse import quote
from typing import Any

import openpyxl
import requests
import xlsxwriter
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.cell import coordinate_to_tuple
from xlsxwriter.utility import xl_range_abs, xl_rowcol_to_cell

from .engine import ModelError, recalculate, validate_inputs

MAX_TEMPLATE_BYTES = 20 * 1024 * 1024
MAX_OUTPUT_CELLS = 15_000
MAX_OUTPUT_STRING = 10_000
MAX_OUTPUT_BYTES = 1_500_000
MAX_SHEETS = 8
MAX_PRESENTATION_SUMMARY = 30
MAX_PRESENTATION_CHARTS = 8
MAX_PRESENTATION_SERIES = 12
MAX_PRESENTATION_SERIES_PER_CHART = 10
MAX_CHART_POINTS = 100
_IDENTIFIER = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,99}$")
_SHA256 = re.compile(r"^[a-fA-F0-9]{64}$")
_CELL = re.compile(r"^[A-Z]{1,3}[1-9][0-9]{0,6}$")


class PreviewError(Exception):
    """A safe-to-return model preview failure."""

    def __init__(self, code: str, reason: str | None = None):
        super().__init__(code)
        self.code = code
        self.reason = reason if reason == "IRR_UNDEFINED" else None


def _supabase_headers(service_key: str) -> dict[str, str]:
    headers = {"apikey": service_key, "Accept": "application/json"}
    # New Supabase secret keys are API keys, not JWT bearer tokens.
    if not service_key.startswith("sb_secret_"):
        headers["Authorization"] = f"Bearer {service_key}"
    return headers


def _validate_identity(calculator_id: Any, version: Any) -> tuple[str, str]:
    if not isinstance(calculator_id, str) or not _IDENTIFIER.fullmatch(calculator_id):
        raise PreviewError("INVALID_REQUEST")
    if not isinstance(version, str) or not _IDENTIFIER.fullmatch(version):
        raise PreviewError("INVALID_REQUEST")
    return calculator_id, version


def _validate_manifest_record(row: Any, calculator_id: str, version: str) -> dict[str, Any]:
    if not isinstance(row, dict) or row.get("approved") is not True:
        # Missing and unapproved candidates intentionally have the same
        # outward result, preventing version review state enumeration.
        raise PreviewError("MODEL_UNAVAILABLE")
    manifest = row.get("manifest")
    if not isinstance(manifest, dict) or manifest.get("id") != calculator_id or manifest.get("version") != version:
        raise PreviewError("MODEL_UNAVAILABLE")
    runtime_sha = manifest.get("runtimeSha256")
    template_sha = row.get("template_sha256")
    path = row.get("template_path")
    if (
        not isinstance(runtime_sha, str)
        or not _SHA256.fullmatch(runtime_sha)
        or not isinstance(template_sha, str)
        or not _SHA256.fullmatch(template_sha)
        or runtime_sha.lower() != template_sha.lower()
        or not isinstance(path, str)
        or path != f"{calculator_id}/{version}/{template_sha.lower()}.xlsx"
    ):
        raise PreviewError("MODEL_UNAVAILABLE")
    return manifest


def load_approved_model(
    calculator_id: str,
    version: str,
    template_path: Path,
    *,
    supabase_url: str,
    service_key: str,
) -> dict[str, Any]:
    """Fetch a fresh approval record and its private immutable template."""
    calculator_id, version = _validate_identity(calculator_id, version)
    if not supabase_url or not service_key:
        raise PreviewError("ENGINE_NOT_CONFIGURED")
    base = supabase_url.rstrip("/")
    headers = _supabase_headers(service_key)
    try:
        response = requests.get(
            f"{base}/rest/v1/calculator_versions",
            params={
                "select": "manifest,template_path,template_sha256,approved",
                "calculator_id": f"eq.{calculator_id}",
                "version": f"eq.{version}",
                "limit": "2",
            },
            headers=headers,
            timeout=(5, 15),
        )
        response.raise_for_status()
        rows = response.json()
    except (requests.RequestException, ValueError):
        raise PreviewError("MODEL_UNAVAILABLE") from None
    if not isinstance(rows, list) or len(rows) != 1:
        raise PreviewError("MODEL_UNAVAILABLE")
    row = rows[0]
    manifest = _validate_manifest_record(row, calculator_id, version)

    storage_path = row["template_path"]
    encoded_path = quote(storage_path, safe="/")
    try:
        with requests.get(
            f"{base}/storage/v1/object/authenticated/model-templates/{encoded_path}",
            headers=headers,
            stream=True,
            timeout=(5, 30),
        ) as download:
            download.raise_for_status()
            content_length = download.headers.get("Content-Length")
            if content_length and int(content_length) > MAX_TEMPLATE_BYTES:
                raise PreviewError("MODEL_UNAVAILABLE")
            total = 0
            template_path.parent.mkdir(parents=True, exist_ok=True)
            with template_path.open("wb") as target:
                for chunk in download.iter_content(64 * 1024):
                    if not chunk:
                        continue
                    total += len(chunk)
                    if total > MAX_TEMPLATE_BYTES:
                        raise PreviewError("MODEL_UNAVAILABLE")
                    target.write(chunk)
    except PreviewError:
        raise
    except (requests.RequestException, OSError, ValueError):
        raise PreviewError("MODEL_UNAVAILABLE") from None

    digest = hashlib.sha256(template_path.read_bytes()).hexdigest()
    if digest != row["template_sha256"].lower() or digest != manifest["runtimeSha256"].lower():
        raise PreviewError("MODEL_UNAVAILABLE")
    return manifest


def _views(manifest: dict[str, Any]) -> list[dict[str, Any]]:
    outputs = manifest.get("outputs")
    views = outputs.get("views") if isinstance(outputs, dict) else None
    if not isinstance(views, list) or not 1 <= len(views) <= MAX_SHEETS:
        raise PreviewError("OUTPUT_MAPPING_INVALID")
    seen_ids: set[str] = set()
    seen_sheets: set[str] = set()
    normalized: list[dict[str, Any]] = []
    for view in views:
        if not isinstance(view, dict):
            raise PreviewError("OUTPUT_MAPPING_INVALID")
        view_id, label, sheet_name, source_range = (
            view.get("id"), view.get("label"), view.get("sheet"), view.get("sourceRange")
        )
        if (
            not isinstance(view_id, str) or not view_id or len(view_id) > 100
            or view_id in seen_ids
            or not isinstance(label, str) or not label or len(label) > 200
            or not isinstance(sheet_name, str) or not sheet_name.endswith(" OUTPUT") or len(sheet_name) > 100
            or sheet_name in seen_sheets
            or not isinstance(source_range, str) or len(source_range) > 100
        ):
            raise PreviewError("OUTPUT_MAPPING_INVALID")
        try:
            min_col, min_row, max_col, max_row = range_boundaries(source_range)
        except (ValueError, TypeError):
            raise PreviewError("OUTPUT_MAPPING_INVALID") from None
        if (
            not all(isinstance(v, int) for v in (min_col, min_row, max_col, max_row))
            or min_col < 1 or min_row < 1 or max_col > 500 or max_row > 1000
            or max_col < min_col or max_row < min_row
            or (max_col - min_col + 1) * (max_row - min_row + 1) > MAX_OUTPUT_CELLS
        ):
            raise PreviewError("OUTPUT_MAPPING_INVALID")
        seen_ids.add(view_id)
        seen_sheets.add(sheet_name)
        normalized.append({
            "id": view_id,
            "label": label,
            "sheet": sheet_name,
            "bounds": (min_col, min_row, max_col, max_row),
        })
    return normalized


def _hidden_columns(sheet: Any) -> set[int]:
    hidden: set[int] = set()
    for key, dimension in sheet.column_dimensions.items():
        if not dimension.hidden:
            continue
        start = dimension.min
        end = dimension.max
        if start is None or end is None:
            match = re.fullmatch(r"[A-Z]+", key)
            if not match:
                continue
            from openpyxl.utils.cell import column_index_from_string
            start = end = column_index_from_string(key)
        hidden.update(range(start, end + 1))
    return hidden


def _json_value(value: Any) -> str | int | float | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date, datetime_time)):
        return value.isoformat()
    if isinstance(value, str):
        if len(value) > MAX_OUTPUT_STRING:
            raise PreviewError("RESULT_INVALID")
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise PreviewError("RESULT_INVALID")
        return value
    raise PreviewError("RESULT_INVALID")


def _input_cells(manifest: dict[str, Any], normalized: dict[str, Any]) -> tuple[str, dict[str, Any]]:
    input_sheet = manifest.get("inputSheet", "Inputs")
    if not isinstance(input_sheet, str) or not input_sheet or len(input_sheet) > 100:
        raise PreviewError("MODEL_UNAVAILABLE")
    cells: dict[str, Any] = {}
    for field in manifest.get("fields", []):
        raw_address = field.get("cell")
        if not isinstance(raw_address, str):
            raise PreviewError("MODEL_UNAVAILABLE")
        if "!" in raw_address:
            sheet_name, raw_address = raw_address.rsplit("!", 1)
            sheet_name = sheet_name.strip("'").replace("''", "'")
            if sheet_name != input_sheet:
                raise PreviewError("MODEL_UNAVAILABLE")
        address = raw_address.replace("$", "").upper()
        if not _CELL.fullmatch(address):
            raise PreviewError("MODEL_UNAVAILABLE")
        try:
            row, column = coordinate_to_tuple(address)
        except ValueError:
            raise PreviewError("MODEL_UNAVAILABLE") from None
        if row > 1_048_576 or column > 16_384 or address in cells:
            raise PreviewError("MODEL_UNAVAILABLE")
        cells[address] = normalized[field["id"]]
    return input_sheet, cells


def normalize_model_inputs(inputs: Any, manifest: dict[str, Any]) -> dict[str, Any]:
    """Validate and normalize field IDs before mapping values to workbook cells.

    Kept separate from recalculation so both the cold path and a persistent
    office-engine adapter can apply the same public-input contract.
    """
    if not isinstance(inputs, dict):
        raise PreviewError("INVALID_REQUEST")
    try:
        normalized = validate_inputs(inputs, manifest)
        fields = manifest.get("fields")
        if not isinstance(fields, list):
            raise ModelError("Invalid model fields")
        by_id = {field.get("id"): field for field in fields if isinstance(field, dict)}
        for field in fields:
            if not isinstance(field, dict) or "dependsOn" not in field:
                continue
            parent_id = field.get("dependsOn")
            allowed_by_parent = field.get("optionsByDependencyValue")
            if not isinstance(parent_id, str) or parent_id not in normalized or not isinstance(allowed_by_parent, dict):
                raise ModelError("Invalid dependent input mapping")
            options = allowed_by_parent.get(normalized[parent_id])
            value = normalized.get(field.get("id"))
            if not isinstance(options, list) or value not in options:
                label = field.get("label") if isinstance(field.get("label"), str) else "Dependent input"
                raise ModelError("Invalid selection: " + label)
        # Every dependency must point at a field rather than being ignored.
        if any(field.get("dependsOn") not in by_id for field in fields if isinstance(field, dict) and "dependsOn" in field):
            raise ModelError("Invalid dependent input mapping")
        return normalized
    except PreviewError:
        raise
    except (ModelError, KeyError, TypeError, ValueError):
        raise PreviewError("INVALID_REQUEST") from None


def serialize_output_sheets(workbook: Any, manifest: dict[str, Any]) -> list[dict[str, Any]]:
    """Return values/number formats from manifest-allowlisted visible ranges."""
    views = _views(manifest)
    result = []
    output_bytes = 0
    for view in views:
        name = view["sheet"]
        if name not in workbook.sheetnames:
            raise PreviewError("OUTPUT_MAPPING_INVALID")
        sheet = workbook[name]
        if sheet.sheet_state != "visible":
            raise PreviewError("OUTPUT_MAPPING_INVALID")
        min_col, min_row, max_col, max_row = view["bounds"]
        hidden_columns = _hidden_columns(sheet)
        rows = []
        cells_written = 0
        output_bytes += len(name.encode("utf-8")) + 32
        for row_index in range(min_row, max_row + 1):
            if sheet.row_dimensions[row_index].hidden:
                continue
            cells = []
            for column_index in range(min_col, max_col + 1):
                if column_index in hidden_columns:
                    continue
                cell = sheet.cell(row=row_index, column=column_index)
                if cell.data_type == "e":
                    raise PreviewError("RESULT_INVALID")
                value = _json_value(cell.value)
                if value is None:
                    continue
                output_bytes += (len(value.encode("utf-8")) if isinstance(value, str) else 32) + 48
                item: dict[str, Any] = {"column": column_index, "value": value}
                number_format = cell.number_format
                if isinstance(number_format, str) and number_format and len(number_format) <= 200:
                    output_bytes += len(number_format.encode("utf-8"))
                    item["format"] = number_format
                cells.append(item)
                cells_written += 1
                if cells_written > MAX_OUTPUT_CELLS or output_bytes > MAX_OUTPUT_BYTES:
                    raise PreviewError("RESULT_INVALID")
            rows.append({"cells": cells})
        result.append({"name": name, "rows": rows})
    return result


def _presentation_specs(manifest: dict[str, Any]) -> dict[str, Any] | None:
    """Validate private summary/chart mappings and return safe cell geometry.

    Presentation metadata is an explicit allowlist. The returned public payload
    never contains workbook formulas, sheet names for chart sources, or source
    ranges; hidden calculation sheets are read only for the mapped chart values.
    """
    raw = manifest.get("presentation")
    if raw is None:
        return None
    if not isinstance(raw, dict):
        raise PreviewError("OUTPUT_MAPPING_INVALID")
    views = _views(manifest)
    by_sheet = {view["sheet"]: view for view in views}
    summary_specs = raw.get("summary", [])
    chart_specs = raw.get("charts", [])
    if (
        not isinstance(summary_specs, list) or len(summary_specs) > MAX_PRESENTATION_SUMMARY
        or not isinstance(chart_specs, list) or len(chart_specs) > MAX_PRESENTATION_CHARTS
        or (not summary_specs and not chart_specs)
    ):
        raise PreviewError("OUTPUT_MAPPING_INVALID")

    def cell_in_view(address: Any, view: dict[str, Any]) -> tuple[int, int]:
        if not isinstance(address, str) or not _CELL.fullmatch(address):
            raise PreviewError("OUTPUT_MAPPING_INVALID")
        try:
            row, column = coordinate_to_tuple(address)
        except ValueError:
            raise PreviewError("OUTPUT_MAPPING_INVALID") from None
        left, top, right, bottom = view["bounds"]
        if not (left <= column <= right and top <= row <= bottom):
            raise PreviewError("OUTPUT_MAPPING_INVALID")
        return row, column

    normalized_summary = []
    seen_summary_cells: set[tuple[str, str]] = set()
    for item in summary_specs:
        if not isinstance(item, dict):
            raise PreviewError("OUTPUT_MAPPING_INVALID")
        section, label, sheet_name, label_cell, value_cell, unit = (
            item.get("section"), item.get("label"), item.get("sheet"),
            item.get("labelCell"), item.get("valueCell"), item.get("unit"),
        )
        if (
            (section is not None and (not isinstance(section, str) or not section or len(section) > 100))
            or not isinstance(label, str) or not label or len(label) > 120
            or sheet_name not in by_sheet
            or (unit is not None and (not isinstance(unit, str) or len(unit) > 100))
        ):
            raise PreviewError("OUTPUT_MAPPING_INVALID")
        view = by_sheet[sheet_name]
        if label_cell is not None:
            cell_in_view(label_cell, view)
        cell_in_view(value_cell, view)
        key = (sheet_name, value_cell)
        if key in seen_summary_cells:
            raise PreviewError("OUTPUT_MAPPING_INVALID")
        seen_summary_cells.add(key)
        normalized_summary.append({
            "section": section,
            "label": label,
            "sheet": sheet_name,
            "labelCell": label_cell,
            "valueCell": value_cell,
            "unit": unit,
        })

    normalized_charts = []
    total_series = 0
    for chart in chart_specs:
        if not isinstance(chart, dict):
            raise PreviewError("OUTPUT_MAPPING_INVALID")
        title, sheet_name, anchor = chart.get("title"), chart.get("sheet"), chart.get("anchor")
        if sheet_name is None:
            dashboard_sheets = [name for name in by_sheet if "dashboard" in name.casefold()]
            if len(dashboard_sheets) == 1:
                sheet_name = dashboard_sheets[0]
        chart_types, groups = chart.get("type"), chart.get("chartGroups")
        if (
            not isinstance(title, str) or not title or len(title) > 120
            or sheet_name not in by_sheet
            or chart.get("categoryRange") is not None
            or chart.get("categoryMode") not in (None, "implicit point index; workbook has no category reference")
        ):
            raise PreviewError("OUTPUT_MAPPING_INVALID")
        if anchor is None:
            anchor_bounds = None
        elif isinstance(anchor, str) and len(anchor) <= 100:
            try:
                anchor_bounds = range_boundaries(anchor)
            except (ValueError, TypeError):
                raise PreviewError("OUTPUT_MAPPING_INVALID") from None
            if not all(isinstance(value, int) and 1 <= value <= 16_384 for value in anchor_bounds):
                raise PreviewError("OUTPUT_MAPPING_INVALID")
        else:
            raise PreviewError("OUTPUT_MAPPING_INVALID")
        if chart.get("kind") in {"line", "column"}:
            kind = chart["kind"]
        elif isinstance(groups, list) and len(groups) == 1 and isinstance(groups[0], dict):
            if chart_types == ["lineChart"] and groups[0].get("type") == "lineChart":
                kind = "line"
            elif (
                chart_types == ["barChart"] and groups[0].get("type") == "barChart"
                and groups[0].get("barDirection") == "col"
            ):
                kind = "column"
            else:
                raise PreviewError("OUTPUT_MAPPING_INVALID")
        else:
            raise PreviewError("OUTPUT_MAPPING_INVALID")
        series_specs = chart.get("series")
        if not isinstance(series_specs, list) or not 1 <= len(series_specs) <= MAX_PRESENTATION_SERIES_PER_CHART:
            raise PreviewError("OUTPUT_MAPPING_INVALID")
        total_series += len(series_specs)
        if total_series > MAX_PRESENTATION_SERIES:
            raise PreviewError("OUTPUT_MAPPING_INVALID")
        normalized_series = []
        lengths = set()
        names = set()
        for series in series_specs:
            if not isinstance(series, dict):
                raise PreviewError("OUTPUT_MAPPING_INVALID")
            name, source_sheet, source_range = (
                series.get("name"), series.get("sheet"), series.get("range")
            )
            if (
                not isinstance(name, str) or not name or len(name) > 100 or name in names
                or not isinstance(source_sheet, str)
                or not (source_sheet in by_sheet or source_sheet.startswith("MODEL "))
                or source_sheet.casefold() in {"ui", "inputs"}
                or not isinstance(source_range, str) or len(source_range) > 100
            ):
                raise PreviewError("OUTPUT_MAPPING_INVALID")
            try:
                bounds = range_boundaries(source_range)
            except (ValueError, TypeError):
                raise PreviewError("OUTPUT_MAPPING_INVALID") from None
            if not all(isinstance(value, int) for value in bounds):
                raise PreviewError("OUTPUT_MAPPING_INVALID")
            left, top, right, bottom = bounds
            length = bottom - top + 1 if left == right else right - left + 1 if top == bottom else 0
            if not 1 <= length <= MAX_CHART_POINTS:
                raise PreviewError("OUTPUT_MAPPING_INVALID")
            normalized_series.append({
                "name": name,
                "sheet": source_sheet,
                "bounds": bounds,
                "length": length,
            })
            lengths.add(length)
            names.add(name)
        if len(lengths) != 1:
            raise PreviewError("OUTPUT_MAPPING_INVALID")
        normalized_charts.append({
            "title": title,
            "kind": kind,
            "sheet": sheet_name,
            "anchor": anchor,
            "anchorBounds": anchor_bounds,
            "pointCount": next(iter(lengths)),
            "series": normalized_series,
        })
    return {"summary": normalized_summary, "charts": normalized_charts}


def serialize_presentation(workbook: Any, manifest: dict[str, Any]) -> dict[str, Any]:
    """Read only mapped dashboard cells and chart series as fresh values."""
    specs = _presentation_specs(manifest)
    if specs is None:
        return {}
    summary = []
    for item in specs["summary"]:
        sheet = workbook[item["sheet"]]
        if sheet.sheet_state != "visible":
            raise PreviewError("OUTPUT_MAPPING_INVALID")
        row, column = coordinate_to_tuple(item["valueCell"])
        if sheet.row_dimensions[row].hidden or column in _hidden_columns(sheet):
            raise PreviewError("OUTPUT_MAPPING_INVALID")
        cell = sheet[item["valueCell"]]
        if cell.data_type == "e":
            raise PreviewError("RESULT_INVALID")
        value = _json_value(cell.value)
        if isinstance(value, str) and len(value) > 500:
            raise PreviewError("RESULT_INVALID")
        result_item = {"label": item["label"], "value": value}
        if item["section"] is not None:
            result_item["section"] = item["section"]
        if isinstance(cell.number_format, str) and len(cell.number_format) <= 200:
            result_item["format"] = cell.number_format
        if item["unit"] is not None:
            result_item["unit"] = item["unit"]
        summary.append(result_item)

    charts = []
    for chart in specs["charts"]:
        chart_series = []
        for item in chart["series"]:
            if item["sheet"] not in workbook.sheetnames:
                raise PreviewError("OUTPUT_MAPPING_INVALID")
            source = workbook[item["sheet"]]
            left, top, right, bottom = item["bounds"]
            raw_values = source.iter_rows(
                min_row=top, max_row=bottom, min_col=left, max_col=right
            )
            values = []
            for row in raw_values:
                for cell in row:
                    if cell.data_type == "e":
                        raise PreviewError("RESULT_INVALID")
                    value = _json_value(cell.value)
                    if value is not None and (
                        isinstance(value, bool) or not isinstance(value, (int, float))
                    ):
                        raise PreviewError("RESULT_INVALID")
                    values.append(value)
            if len(values) != item["length"]:
                raise PreviewError("RESULT_INVALID")
            chart_series.append({"name": item["name"], "values": values})
        charts.append({
            "title": chart["title"],
            "kind": chart["kind"],
            "categories": list(range(1, chart["pointCount"] + 1)),
            "series": chart_series,
        })
    return {"summary": summary, "charts": charts}


def scan_formula_errors(path: str | Path, value_workbook: Any) -> None:
    """Fail if recalculation left an error in any formula, including hidden sheets.

    Output-range validation alone is insufficient: IFERROR can turn a broken
    internal dependency into a plausible zero. Formula text is inspected only
    to identify formula cells and is never returned or logged.
    """
    try:
        formulas = openpyxl.load_workbook(path, data_only=False, read_only=True)
    except Exception:
        raise PreviewError("CALCULATION_FAILED") from None
    try:
        error_formulas = []
        for sheet_name in formulas.sheetnames:
            if sheet_name not in value_workbook.sheetnames:
                raise PreviewError("CALCULATION_FAILED")
            formula_sheet = formulas[sheet_name]
            result_sheet = value_workbook[sheet_name]
            for row in formula_sheet.iter_rows():
                for formula_cell in row:
                    if formula_cell.data_type != "f":
                        continue
                    result_cell = result_sheet[formula_cell.coordinate]
                    if result_cell.data_type == "e":
                        error_formulas.append(formula_cell.value)
        if error_formulas:
            from .formula_diagnostics import formula_failure_reason
            reason = formula_failure_reason(error_formulas)
            raise PreviewError("RESULT_INVALID" if reason else "CALCULATION_FAILED", reason=reason)
    except PreviewError:
        raise
    except Exception:
        raise PreviewError("CALCULATION_FAILED") from None
    finally:
        formulas.close()


def evaluate_template(
    template: str | Path,
    manifest: dict[str, Any],
    inputs: Any,
    *,
    timeout: float = 60,
    workdir: str | Path | None = None,
) -> dict[str, Any]:
    """Recalculate one transient input state and return only allowlisted values."""
    try:
        _views(manifest)
        normalized = normalize_model_inputs(inputs, manifest)
        input_sheet, cells = _input_cells(manifest, normalized)

        started = time.monotonic()
        with tempfile.TemporaryDirectory(prefix="geoflow-preview-", dir=workdir) as temp:
            try:
                calculated_path, workbook = recalculate(
                    template,
                    cells,
                    Path(temp) / "calculation",
                    timeout=timeout,
                    input_sheet=input_sheet,
                )
            except Exception:
                # Do not return engine diagnostics, formulas, or host paths.
                raise PreviewError("CALCULATION_FAILED") from None
            try:
                scan_formula_errors(calculated_path, workbook)
                output_sheets = serialize_output_sheets(workbook, manifest)
                presentation = serialize_presentation(workbook, manifest)
            finally:
                workbook.close()
        duration = round(time.monotonic() - started, 3)
        return {
            "calculatorId": manifest["id"],
            "calculatorVersion": manifest["version"],
            "outputSheets": output_sheets,
            **presentation,
            "durationSeconds": duration,
        }
    except PreviewError:
        raise
    except (KeyError, TypeError, ValueError):
        raise PreviewError("INVALID_REQUEST") from None


def export_preview_results(
    workbook: Any,
    manifest: dict[str, Any],
    target: str | Path,
    *,
    presentation: dict[str, Any] | None = None,
) -> bool:
    """Write a fresh values-only XLSX containing only mapped visible ranges."""
    views = _views(manifest)
    if presentation is None:
        presentation = serialize_presentation(workbook, manifest)
    presentation_specs = _presentation_specs(manifest)
    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    book = xlsxwriter.Workbook(str(target), {
        "strings_to_formulas": False,
        "strings_to_urls": False,
        "nan_inf_to_errors": False,
    })
    format_cache: dict[str, Any] = {}
    allowed_cells: dict[str, set[str]] = {view["sheet"]: set() for view in views}
    chart_formula_refs: list[str] = []
    layouts = {}
    try:
        book.set_properties({
            "title": "Whittaker Geo calculation results",
            "author": "Whittaker Geo",
            "comments": "Values-only mapped model outputs.",
        })
        for view in views:
            source = workbook[view["sheet"]]
            if source.sheet_state != "visible":
                raise PreviewError("OUTPUT_MAPPING_INVALID")
            destination = book.add_worksheet(view["sheet"])
            destination.hide_gridlines(2)
            from .export_layout import write_compact_output
            layouts[view["sheet"]] = write_compact_output(book, source, destination, view,
                _hidden_columns(source), allowed_cells[view["sheet"]], format_cache, _json_value)

        if presentation_specs is not None:
            next_helper_column: dict[str, int] = {}
            next_chart_row: dict[str, int] = {}
            for view in views:
                geometry = layouts[view["sheet"]]
                next_helper_column[view["sheet"]] = geometry["columnCount"] + 1
                next_chart_row[view["sheet"]] = geometry["rowCount"] + 2
            for spec, chart_result in zip(presentation_specs["charts"], presentation["charts"]):
                destination = book.get_worksheet_by_name(spec["sheet"])
                if destination is None:
                    raise PreviewError("OUTPUT_MAPPING_INVALID")
                categories = chart_result["categories"]
                start_column = next_helper_column[spec["sheet"]]
                start_row = 0
                category_column = start_column
                destination.write_string(start_row, category_column, "Point")
                for point, category in enumerate(categories, start=1):
                    destination.write_number(start_row + point, category_column, category)
                allowed_cells[spec["sheet"]].add(
                    f"{xl_rowcol_to_cell(start_row, category_column)}"
                )
                for point in range(1, len(categories) + 1):
                    allowed_cells[spec["sheet"]].add(
                        f"{xl_rowcol_to_cell(start_row + point, category_column)}"
                    )

                chart = book.add_chart({"type": chart_result["kind"]})
                chart.set_title({"name": chart_result["title"]})
                chart.set_style(10)
                chart.show_hidden_data()
                if len(chart_result["series"]) > 1:
                    chart.set_legend({"position": "bottom"})
                for series_index, series in enumerate(chart_result["series"], start=1):
                    column = start_column + series_index
                    destination.write_string(start_row, column, series["name"])
                    for point, value in enumerate(series["values"], start=1):
                        if value is None:
                            destination.write_blank(start_row + point, column, None)
                        elif isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value):
                            destination.write_number(start_row + point, column, value)
                        else:
                            raise PreviewError("RESULT_INVALID")
                    allowed_cells[spec["sheet"]].add(xl_rowcol_to_cell(start_row, column))
                    for point in range(1, len(categories) + 1):
                        allowed_cells[spec["sheet"]].add(
                            xl_rowcol_to_cell(start_row + point, column)
                        )
                    category_ref = (
                        f"'{spec['sheet']}'!{xl_range_abs(start_row + 1, category_column, start_row + len(categories), category_column)}"
                    )
                    value_ref = (
                        f"'{spec['sheet']}'!{xl_range_abs(start_row + 1, column, start_row + len(categories), column)}"
                    )
                    chart_formula_refs.extend((category_ref, value_ref))
                    chart.add_series({
                        "name": series["name"],
                        "categories": [spec["sheet"], start_row + 1, category_column,
                                       start_row + len(categories), category_column],
                        "values": [spec["sheet"], start_row + 1, column,
                                   start_row + len(categories), column],
                    })

                if spec["anchorBounds"] is not None:
                    left, top, right, bottom = spec["anchorBounds"]
                    width = max(420, (right - left + 1) * 64)
                    height = max(260, (bottom - top + 1) * 20)
                    chart.set_size({"width": width, "height": height})
                    geometry = layouts[spec["sheet"]]
                    chart_row = sum(r < top for r in geometry["rows"])
                    chart_column = sum(c < left for c in geometry["columns"])
                    anchor_cell = xl_rowcol_to_cell(chart_row, chart_column)
                else:
                    output_view = next(view for view in views if view["sheet"] == spec["sheet"])
                    chart_left = 0
                    anchor_cell = xl_rowcol_to_cell(next_chart_row[spec["sheet"]], chart_left)
                    chart.set_size({"width": 900, "height": 320})
                    next_chart_row[spec["sheet"]] += 17
                destination.insert_chart(anchor_cell, chart)
                final_column = start_column + len(chart_result["series"])
                destination.set_column(start_column, final_column, 14, None, {"hidden": True})
                next_helper_column[spec["sheet"]] = final_column + 2
        book.close()
    except Exception:
        try:
            book.close()
        except Exception:
            pass
        target.unlink(missing_ok=True)
        raise
    return inspect_preview_export(
        target,
        [view["sheet"] for view in views],
        expected_chart_count=len(presentation.get("charts", [])),
        allowed_cells=allowed_cells,
        expected_chart_refs=chart_formula_refs,
    )


def inspect_preview_export(
    path: str | Path,
    expected_sheets: list[str],
    *,
    expected_chart_count: int = 0,
    allowed_cells: dict[str, set[str]] | None = None,
    expected_chart_refs: list[str] | None = None,
) -> bool:
    """Verify a saved preview export has only visible allowlisted values."""
    import zipfile
    import xml.etree.ElementTree as ET

    ns = {
        "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    }
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        book = ET.fromstring(archive.read("xl/workbook.xml"))
        sheets = book.findall("s:sheets/s:sheet", ns)
        if [sheet.attrib["name"] for sheet in sheets] != expected_sheets:
            raise PreviewError("EXPORT_INVALID")
        if any(sheet.attrib.get("state", "visible") != "visible" for sheet in sheets):
            raise PreviewError("EXPORT_INVALID")
        if any(any(token in name for token in ("externalLinks", "embeddings", "vbaProject", "comments", "connections")) for name in names):
            raise PreviewError("EXPORT_INVALID")
        chart_files = sorted(
            name for name in names
            if name.startswith("xl/charts/chart") and name.endswith(".xml")
        )
        if len(chart_files) != expected_chart_count:
            raise PreviewError("EXPORT_INVALID")
        actual_chart_refs = []
        for chart_file in chart_files:
            chart_root = ET.fromstring(archive.read(chart_file))
            for formula in chart_root.findall(".//c:f", ns):
                if not formula.text or "[" in formula.text:
                    raise PreviewError("EXPORT_INVALID")
                actual_chart_refs.append(formula.text)
        if expected_chart_refs is not None and sorted(actual_chart_refs) != sorted(expected_chart_refs):
            raise PreviewError("EXPORT_INVALID")
        for name in names:
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
                root = ET.fromstring(archive.read(name))
                if root.findall(".//s:f", ns):
                    raise PreviewError("EXPORT_INVALID")
        if allowed_cells is not None:
            output = openpyxl.load_workbook(path, data_only=False, read_only=True)
            try:
                if output.sheetnames != expected_sheets:
                    raise PreviewError("EXPORT_INVALID")
                for sheet_name in expected_sheets:
                    sheet = output[sheet_name]
                    allowed = allowed_cells.get(sheet_name, set())
                    for row in sheet.iter_rows():
                        for cell in row:
                            if cell.value is not None and cell.coordinate not in allowed:
                                raise PreviewError("EXPORT_INVALID")
            finally:
                output.close()
    return True


def run_saved_model(
    template: str | Path,
    manifest: dict[str, Any],
    inputs: Any,
    output: str | Path,
    *,
    timeout: float = 1800,
) -> dict[str, Any]:
    """Durable-run adapter for the mapped UI/OUTPUT workbook family."""
    expected = manifest.get("runtimeSha256")
    if not isinstance(expected, str) or not _SHA256.fullmatch(expected):
        raise PreviewError("MODEL_UNAVAILABLE")
    if hashlib.sha256(Path(template).read_bytes()).hexdigest() != expected.lower():
        raise PreviewError("MODEL_UNAVAILABLE")
    try:
        _views(manifest)
        normalized = normalize_model_inputs(inputs, manifest)
        input_sheet, cells = _input_cells(manifest, normalized)
        started = time.monotonic()
        with tempfile.TemporaryDirectory(prefix="geoflow-run-") as temp:
            try:
                calculated_path, workbook = recalculate(
                    template, cells, Path(temp) / "calculation", timeout=timeout, input_sheet=input_sheet
                )
            except Exception:
                raise PreviewError("CALCULATION_FAILED") from None
            try:
                scan_formula_errors(calculated_path, workbook)
                output_sheets = serialize_output_sheets(workbook, manifest)
                presentation = serialize_presentation(workbook, manifest)
                export_preview_results(workbook, manifest, output, presentation=presentation)
            finally:
                workbook.close()
        return {
            "calculatorId": manifest["id"],
            "calculatorVersion": manifest["version"],
            "outputSheets": output_sheets,
            **presentation,
            "durationSeconds": round(time.monotonic() - started, 3),
        }
    except PreviewError:
        raise
    except (KeyError, TypeError, ValueError, OSError):
        raise PreviewError("INVALID_REQUEST") from None


def review_candidate(
    manifest_path: str | Path,
    template_path: str | Path,
    inputs_path: str | Path,
) -> dict[str, Any]:
    """Owner-local numerical review; does not require approval or contact DB."""
    manifest = json.loads(Path(manifest_path).read_text(encoding="utf-8"))
    template = Path(template_path)
    expected = manifest.get("runtimeSha256")
    if not isinstance(expected, str) or not _SHA256.fullmatch(expected):
        raise PreviewError("MODEL_UNAVAILABLE")
    if hashlib.sha256(template.read_bytes()).hexdigest() != expected.lower():
        raise PreviewError("MODEL_UNAVAILABLE")
    inputs = json.loads(Path(inputs_path).read_text(encoding="utf-8"))
    return evaluate_template(template, manifest, inputs)


def review_main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Locally review a private model candidate; no HTTP server is started.")
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--inputs", type=Path, required=True, help="JSON keyed by private manifest field IDs.")
    args = parser.parse_args(argv)
    result = review_candidate(args.manifest, args.template, args.inputs)
    print(json.dumps(result, ensure_ascii=False, allow_nan=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(review_main())
