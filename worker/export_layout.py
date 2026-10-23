"""Presentation-only output layout. Numeric values and formats stay unchanged."""
import math
import re
from openpyxl.utils import get_column_letter
from xlsxwriter.utility import xl_rowcol_to_cell


def write_compact_output(book, source, destination, view, hidden_columns, allowed, formats, value_encoder):
    left, top, right, bottom = view["bounds"]
    rows = [r for r in range(top, bottom + 1) if not source.row_dimensions[r].hidden]
    columns = [c for c in range(left, right + 1) if c not in hidden_columns]
    populated = [(r, c) for r in rows for c in columns if source.cell(r, c).value is not None]
    if populated:
        first_row, last_row = min(r for r, _ in populated), max(r for r, _ in populated)
        first_column, last_column = min(c for _, c in populated), max(c for _, c in populated)
        rows = [r for r in rows if first_row <= r <= last_row]
        columns = [c for c in columns if first_column <= c <= last_column]
    row_map = {r: i for i, r in enumerate(rows)}
    column_map = {c: i for i, c in enumerate(columns)}
    year_columns = {c for c in columns if any(isinstance(source.cell(r, c).value, str)
        and re.fullmatch(r"Year(?:\s*\((?:Project|Qi)\))?", source.cell(r, c).value.strip(), re.I) for r in rows)}
    first_data = next((i for i, r in enumerate(rows) if year_columns and any(
        isinstance(source.cell(r, c).value, (int, float)) for c in year_columns)), None)
    widths = {}
    for c, output_column in column_map.items():
        dimension = source.column_dimensions.get(get_column_letter(c))
        width = 9 if c in year_columns else max(20, min(dimension.width if dimension and dimension.width else 23, 30))
        widths[c] = width
        destination.set_column(output_column, output_column, width)
    for r, output_row in row_map.items():
        values = [(c, source.cell(r, c).value) for c in columns if source.cell(r, c).value is not None]
        text_only = bool(values) and all(isinstance(value, str) for _, value in values)
        footnote = text_only and all(str(value).lstrip().startswith("*") for _, value in values)
        header = not footnote and (output_row == 0 or text_only and (first_data is None or output_row < first_data))
        section = header and (output_row == 0 or len(values) <= 3 or any(values[i][0] - values[i-1][0] > 1 for i in range(1, len(values))))
        height = max((min(72, 15 * max(1, math.ceil(len(str(value)) / max(8, widths[c] - 2))))
                      for c, value in values if isinstance(value, str)), default=18)
        destination.set_row(output_row, 28 if section else 20 if footnote else max(20, height + 6))
        for c, output_column in column_map.items():
            cell = source.cell(r, c)
            value = value_encoder(cell.value)
            if cell.data_type == "e" or isinstance(value, float) and not math.isfinite(value):
                raise ValueError("Invalid output value")
            properties = {"font_name": "Calibri", "font_size": 11, "border": 1,
                          "border_color": "#D5DFE2", "valign": "vcenter", "num_format": cell.number_format,
                          "text_wrap": isinstance(value, str)}
            if header:
                properties.update(bold=True, text_wrap=True, bg_color="#002D43" if section else "#E8EDE0",
                                  font_color="#FFFFFF" if section else "#002D43")
            elif output_row % 2:
                properties["bg_color"] = "#F5F7F3"
            key = tuple(sorted(properties.items()))
            if key not in formats:
                formats[key] = book.add_format(properties)
            fmt = formats[key]
            allowed.add(xl_rowcol_to_cell(output_row, output_column))
            if value is None:
                destination.write_blank(output_row, output_column, None, fmt)
            elif isinstance(value, str):
                destination.write_string(output_row, output_column, value, fmt)
            elif isinstance(value, (int, float)) and not isinstance(value, bool):
                destination.write_number(output_row, output_column, value, fmt)
            else:
                raise ValueError("Invalid output value")
        # Span header labels only across blank cells; no values are overwritten.
        if section:
            for i, (c, value) in enumerate(values):
                if not isinstance(value, str):
                    continue
                start = column_map[c]
                end = column_map[values[i + 1][0]] - 1 if i + 1 < len(values) else len(columns) - 1
                if end > start:
                    destination.merge_range(output_row, start, output_row, end, value,
                        book.add_format({"font_name": "Calibri", "font_size": 11, "bold": True,
                          "font_color": "#FFFFFF", "bg_color": "#002D43", "border": 1,
                          "border_color": "#D5DFE2", "text_wrap": True, "valign": "vcenter",
                          "num_format": source.cell(r, c).number_format}))
    destination.freeze_panes(first_data or 0, 1 if year_columns else 0)
    if first_data:
        destination.repeat_rows(0, first_data - 1)
    destination.set_landscape()
    destination.fit_to_pages(1, 0)
    destination.set_zoom(85)
    destination.set_tab_color("#002D43")
    destination.set_selection(0, 0, 0, 0)
    destination.set_top_left_cell(0, 0)
    return {"rows": row_map, "columns": column_map, "rowCount": len(rows), "columnCount": len(columns)}
