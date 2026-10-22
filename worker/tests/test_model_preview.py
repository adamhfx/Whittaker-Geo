import json
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
import pytest

from worker.engine import ModelError, validate_inputs
from worker.model_preview import (
    PreviewError,
    _presentation_specs,
    evaluate_template,
    inspect_preview_export,
    normalize_model_inputs,
    scan_formula_errors,
    serialize_presentation,
    serialize_output_sheets,
    run_saved_model,
)


def manifest():
    return {
        "id": "whittaker-govrev",
        "version": "1.0.0-candidate.1",
        "inputSheet": "UI",
        "fields": [
            {
                "id": "rate",
                "cell": "B2",
                "label": "Oil price",
                "type": "number",
                "required": True,
                "min": 0,
            },
            {
                "id": "facility",
                "cell": "B3",
                "label": "Facility year",
                "type": "milestone",
                "required": False,
            },
        ],
        "outputs": {
            "views": [
                {
                    "id": "summary",
                    "label": "Summary",
                    "sheet": "Executive OUTPUT",
                    "sourceRange": "A1:D3",
                }
            ]
        },
    }


def presentation_manifest():
    model = manifest()
    model["outputs"]["views"] = [
        {
            "id": "dashboard", "label": "Dashboard", "sheet": "Executive OUTPUT",
            "sourceRange": "A1:D3",
        },
        {
            "id": "annual", "label": "Annual results", "sheet": "Annual OUTPUT",
            "sourceRange": "A1:B2",
        },
    ]
    implicit_categories = "implicit point index; workbook has no category reference"
    model["presentation"] = {
        "summary": [{
            "section": "Investor", "label": "NPV", "labelCell": "A1",
            "valueCell": "B1", "sheet": "Executive OUTPUT", "unit": "MM USD",
        }],
        "charts": [
            {
                "title": "Type Well Gas Production", "sheet": "Executive OUTPUT",
                "anchor": "A4:C12", "type": ["lineChart"],
                "chartGroups": [{"type": "lineChart", "barDirection": None, "grouping": "standard"}],
                "categorySheet": None, "categoryRange": None, "categoryMode": implicit_categories,
                "series": [{"name": "Gas Production (MM m3/year)", "sheet": "MODEL Dashboard", "range": "A1:A3"}],
            },
            {
                "title": "Annual Field Production", "sheet": "Executive OUTPUT",
                "anchor": "D4:F12", "type": ["barChart"],
                "chartGroups": [{"type": "barChart", "barDirection": "col", "grouping": "clustered"}],
                "categorySheet": None, "categoryRange": None, "categoryMode": implicit_categories,
                "series": [
                    {"name": "Oil / Liquids (MM boe/year)", "sheet": "Executive OUTPUT", "range": "D1:D3"},
                    {"name": "Gas Equivalent (MM boe/year)", "sheet": "MODEL Dashboard", "range": "B1:B3"},
                ],
            },
            {
                "title": "Cash Flow and Capex Summary", "sheet": "Executive OUTPUT",
                "anchor": "G4:I12", "type": ["barChart"],
                "chartGroups": [{"type": "barChart", "barDirection": "col", "grouping": "clustered"}],
                "categorySheet": None, "categoryRange": None, "categoryMode": implicit_categories,
                "series": [
                    {"name": "Investor CF (MM USD/year)", "sheet": "MODEL Dashboard", "range": "C1:C3"},
                    {"name": "State CF (MM USD/year)", "sheet": "MODEL Dashboard", "range": "D1:D3"},
                    {"name": "Project Capex (MM USD/year)", "sheet": "MODEL Dashboard", "range": "E1:E3"},
                ],
            },
        ],
    }
    return model


def test_presentation_mappings_enforce_client_summary_chart_limits():
    model = presentation_manifest()
    model["presentation"]["summary"] *= 31
    with pytest.raises(PreviewError, match="OUTPUT_MAPPING_INVALID"):
        _presentation_specs(model)

    model = presentation_manifest()
    model["presentation"]["summary"][0]["label"] = "L" * 121
    with pytest.raises(PreviewError, match="OUTPUT_MAPPING_INVALID"):
        _presentation_specs(model)

    model = presentation_manifest()
    model["presentation"]["charts"][0]["series"] *= 11
    with pytest.raises(PreviewError, match="OUTPUT_MAPPING_INVALID"):
        _presentation_specs(model)

    model = presentation_manifest()
    model["presentation"]["charts"][0]["title"] = "T" * 121
    with pytest.raises(PreviewError, match="OUTPUT_MAPPING_INVALID"):
        _presentation_specs(model)

    model = presentation_manifest()
    model["presentation"]["charts"][0]["series"][0]["name"] = "N" * 101
    with pytest.raises(PreviewError, match="OUTPUT_MAPPING_INVALID"):
        _presentation_specs(model)

    model = presentation_manifest()
    model["presentation"]["charts"][0]["series"][0]["range"] = "A1:A101"
    with pytest.raises(PreviewError, match="OUTPUT_MAPPING_INVALID"):
        _presentation_specs(model)


def test_preview_maps_input_ids_to_ui_cells_and_returns_only_visible_values(tmp_path, monkeypatch):
    source = tmp_path / "synthetic.xlsx"
    book = openpyxl.Workbook()
    ui = book.active
    ui.title = "UI"
    ui["B2"] = 10
    ui["B3"] = "X"
    output = book.create_sheet("Executive OUTPUT")
    output["A1"] = "NPV"
    output["B1"] = 125.5
    output["B1"].number_format = '$#,##0.0'
    output["C1"] = "Hidden column secret"
    output["D1"] = "=1+2"
    output["A2"] = "Hidden row secret"
    output["B2"] = 999
    output["A3"] = "Visible row"
    output["D3"] = "inside range"
    output["E1"] = "outside allowlist"
    output.column_dimensions["C"].hidden = True
    output.row_dimensions[2].hidden = True
    book.save(source)

    seen = {}

    def fake_recalculate(template, values, folder, timeout=90, input_sheet="Inputs"):
        seen.update(values=values, input_sheet=input_sheet, timeout=timeout)
        return template, openpyxl.load_workbook(template, data_only=True)

    monkeypatch.setattr("worker.model_preview.recalculate", fake_recalculate)
    result = evaluate_template(
        source,
        manifest(),
        {"rate": 11, "facility": "X"},
        timeout=13,
    )

    assert seen == {"values": {"B2": 11, "B3": "X"}, "input_sheet": "UI", "timeout": 13}
    assert result["calculatorId"] == "whittaker-govrev"
    assert result["calculatorVersion"] == "1.0.0-candidate.1"
    assert [row["cells"] for row in result["outputSheets"][0]["rows"]] == [
        [
            {"column": 1, "value": "NPV", "format": "General"},
            {"column": 2, "value": 125.5, "format": "$#,##0.0"},
        ],
        [
            {"column": 1, "value": "Visible row", "format": "General"},
            {"column": 4, "value": "inside range", "format": "General"},
        ],
    ]
    serialized = json.dumps(result)
    assert "1+2" not in serialized
    assert "Hidden column secret" not in serialized
    assert "Hidden row secret" not in serialized
    assert "outside allowlist" not in serialized
    assert "metrics" not in result and "annual" not in result


def test_preview_rejects_unknown_or_hidden_output_mapping(tmp_path):
    book = openpyxl.Workbook()
    book.active.title = "Allowed"
    with pytest.raises(PreviewError, match="OUTPUT_MAPPING_INVALID"):
        serialize_output_sheets(book, {
            "outputs": {"views": [{
                "id": "unexpected", "label": "Unexpected",
                "sheet": "Private Internal", "sourceRange": "A1:B2",
            }]}
        })

    book["Allowed"].sheet_state = "hidden"
    with pytest.raises(PreviewError, match="OUTPUT_MAPPING_INVALID"):
        serialize_output_sheets(book, {
            "outputs": {"views": [{
                "id": "hidden", "label": "Hidden",
                "sheet": "Allowed", "sourceRange": "A1:B2",
            }]}
        })


def test_preview_rejects_ranges_that_are_not_explicit_bounded_cells():
    invalid = manifest()
    invalid["outputs"]["views"][0]["sourceRange"] = "A:A"
    with pytest.raises(PreviewError, match="OUTPUT_MAPPING_INVALID"):
        serialize_output_sheets(openpyxl.Workbook(), invalid)


def _add_formula_error_caches(path, errors):
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    ns = {"m": namespace, "r": rel_ns, "p": package_rel_ns}
    with zipfile.ZipFile(path) as source:
        workbook = ET.fromstring(source.read("xl/workbook.xml"))
        rels = ET.fromstring(source.read("xl/_rels/workbook.xml.rels"))
        relation_targets = {rel.attrib["Id"]: rel.attrib["Target"] for rel in rels}
        sheet_paths = {}
        for sheet in workbook.findall("m:sheets/m:sheet", ns):
            target = relation_targets[sheet.attrib[f"{{{rel_ns}}}id"]]
            sheet_paths[sheet.attrib["name"]] = target.lstrip("/") if target.startswith("/") else "xl/" + target
        rewritten = {}
        for sheet_name, cell_values in errors.items():
            xml = ET.fromstring(source.read(sheet_paths[sheet_name]))
            cells = {cell.attrib["r"]: cell for cell in xml.findall(".//m:sheetData/m:row/m:c", ns)}
            for address, error in cell_values.items():
                cell = cells[address]
                value = cell.find("m:v", ns)
                if value is None:
                    value = ET.SubElement(cell, f"{{{namespace}}}v")
                value.text = error
                if error.startswith("#"):
                    cell.set("t", "e")
                else:
                    cell.attrib.pop("t", None)
            rewritten[sheet_paths[sheet_name]] = ET.tostring(xml, encoding="utf-8", xml_declaration=True)
        with zipfile.ZipFile(str(path) + ".tmp", "w", zipfile.ZIP_DEFLATED) as target:
            for name in source.namelist():
                target.writestr(name, rewritten.get(name, source.read(name)))
    Path(str(path) + ".tmp").replace(path)


def test_internal_formula_errors_fail_even_when_iferror_hides_them(tmp_path):
    source = tmp_path / "formula-errors.xlsx"
    book = openpyxl.Workbook()
    ui = book.active
    ui.title = "UI"
    ui["B2"] = 10
    ui["B3"] = "X"
    internal = book.create_sheet("Internal")
    internal.sheet_state = "hidden"
    internal["A1"] = "=1/0"
    output = book.create_sheet("Executive OUTPUT")
    output["A1"] = "NPV"
    output["B1"] = "=IFERROR(Internal!A1,0)"
    book.save(source)
    _add_formula_error_caches(source, {"Internal": {"A1": "#DIV/0!"}, "Executive OUTPUT": {"B1": "0"}})

    values = openpyxl.load_workbook(source, data_only=True)
    assert values["Executive OUTPUT"]["B1"].value == 0
    with pytest.raises(PreviewError, match="CALCULATION_FAILED"):
        scan_formula_errors(source, values)
    values.close()


def test_saved_model_dispatch_exports_only_mapped_visible_values(tmp_path, monkeypatch):
    source = tmp_path / "synthetic.xlsx"
    book = openpyxl.Workbook()
    ui = book.active
    ui.title = "UI"
    ui["B2"] = 10
    ui["B3"] = "X"
    output = book.create_sheet("Executive OUTPUT")
    output["A1"] = "NPV"
    output["B1"] = 125.5
    output["E1"] = "Private value outside map"
    book.save(source)
    runtime = dict(manifest(), runtimeSha256=__import__("hashlib").sha256(source.read_bytes()).hexdigest())
    seen = {}

    def fake_recalculate(template, values, folder, timeout=90, input_sheet="Inputs"):
        seen.update(values=values, input_sheet=input_sheet)
        return template, openpyxl.load_workbook(template, data_only=True)

    monkeypatch.setattr("worker.model_preview.recalculate", fake_recalculate)
    export = tmp_path / "results.xlsx"
    result = run_saved_model(source, runtime, {"rate": 12, "facility": "X"}, export)

    assert seen == {"values": {"B2": 12, "B3": "X"}, "input_sheet": "UI"}
    assert result["calculatorId"] == "whittaker-govrev"
    assert result["calculatorVersion"] == "1.0.0-candidate.1"
    assert [sheet["name"] for sheet in result["outputSheets"]] == ["Executive OUTPUT"]
    assert "metrics" not in result and "annual" not in result
    assert inspect_preview_export(export, ["Executive OUTPUT"])
    exported = openpyxl.load_workbook(export, data_only=False)
    assert exported.sheetnames == ["Executive OUTPUT"]
    assert exported["Executive OUTPUT"]["B1"].value == 125.5
    assert exported["Executive OUTPUT"]["E1"].value is None
    assert not any(cell.data_type == "f" for row in exported["Executive OUTPUT"] for cell in row)
    exported.close()


def test_saved_presentation_rebuilds_charts_and_exports_only_allowlisted_values(tmp_path, monkeypatch):
    source = tmp_path / "presentation.xlsx"
    book = openpyxl.Workbook()
    ui = book.active
    ui.title = "UI"
    ui["B2"] = 10
    ui["B3"] = "X"
    dashboard = book.create_sheet("Executive OUTPUT")
    dashboard["A1"] = "NPV"
    dashboard["B1"] = 125.5
    dashboard["B1"].number_format = '$#,##0.0'
    dashboard["D1"], dashboard["D2"], dashboard["D3"] = 1, 2, 3
    dashboard["E1"] = "Outside mapped range"
    annual = book.create_sheet("Annual OUTPUT")
    annual["A1"], annual["B1"] = "Year", 2026
    annual["A2"], annual["B2"] = "Revenue", 5.5
    internal = book.create_sheet("MODEL Dashboard")
    internal.sheet_state = "hidden"
    for column, values in enumerate(([10, 20, 30], [1, 2, 3], [4, 5, 6], [7, 8, 9], [11, 12, 13]), start=1):
        for row, value in enumerate(values, start=1):
            internal.cell(row=row, column=column, value=value)
    internal["Z10"] = "Hidden private sentinel"
    book.save(source)
    model = presentation_manifest()
    model["runtimeSha256"] = __import__("hashlib").sha256(source.read_bytes()).hexdigest()

    def fake_recalculate(template, values, folder, timeout=90, input_sheet="Inputs"):
        assert input_sheet == "UI"
        return template, openpyxl.load_workbook(template, data_only=True)

    monkeypatch.setattr("worker.model_preview.recalculate", fake_recalculate)
    export = tmp_path / "presentation-results.xlsx"
    result = run_saved_model(source, model, {"rate": 12, "facility": "X"}, export)

    assert result["summary"] == [{
        "label": "NPV", "value": 125.5, "section": "Investor",
        "format": '$#,##0.0', "unit": "MM USD",
    }]
    assert [chart["kind"] for chart in result["charts"]] == ["line", "column", "column"]
    assert [chart["categories"] for chart in result["charts"]] == [[1, 2, 3]] * 3
    assert [len(chart["series"]) for chart in result["charts"]] == [1, 2, 3]
    serialized = json.dumps({"summary": result["summary"], "charts": result["charts"]})
    assert "MODEL Dashboard" not in serialized and "A1:A3" not in serialized

    assert inspect_preview_export(export, ["Executive OUTPUT", "Annual OUTPUT"], expected_chart_count=3)
    exported = openpyxl.load_workbook(export, data_only=False)
    try:
        assert exported.sheetnames == ["Executive OUTPUT", "Annual OUTPUT"]
        assert len(exported["Executive OUTPUT"]._charts) == 3
        assert exported["Executive OUTPUT"]["E1"].value is None
        assert exported["Executive OUTPUT"]["F2"].value == 1
        assert exported["Executive OUTPUT"]["G2"].value == 10
        assert exported["Executive OUTPUT"].column_dimensions["F"].hidden
        assert "MODEL Dashboard" not in exported.sheetnames
        assert all(cell.data_type != "f" for sheet in exported for row in sheet.iter_rows() for cell in row)
    finally:
        exported.close()
    with zipfile.ZipFile(export) as archive:
        package_text = b"\n".join(archive.read(name) for name in archive.namelist())
    assert b"Hidden private sentinel" not in package_text
    assert b"MODEL Dashboard" not in package_text


def test_field_validation_does_not_require_a_milestone_when_model_has_none():
    manifest_without_milestone = {
        "fields": [
            {"id": "rate", "label": "Rate", "type": "number", "required": True, "min": 0}
        ]
    }
    assert validate_inputs({"rate": 0}, manifest_without_milestone) == {"rate": 0}


def test_preview_input_validation_rejects_unmapped_field():
    with pytest.raises(PreviewError, match="INVALID_REQUEST"):
        evaluate_template(
            "missing.xlsx",
            manifest(),
            {"rate": 11, "facility": "X", "extra": 1},
        )


def test_dependent_options_are_enforced_by_shared_normalizer():
    m = {
        "fields": [
            {"id": "fluid", "label": "Fluid", "type": "select", "required": True, "options": ["Oil", "NG"]},
            {
                "id": "unit", "label": "Volume unit", "type": "select", "required": True,
                "options": ["barrels(oil)", "m3(gas)"], "dependsOn": "fluid",
                "optionsByDependencyValue": {"Oil": ["barrels(oil)"], "NG": ["m3(gas)"]},
            },
        ]
    }
    assert normalize_model_inputs({"fluid": "Oil", "unit": "barrels(oil)"}, m) == {
        "fluid": "Oil", "unit": "barrels(oil)"
    }
    with pytest.raises(PreviewError, match="INVALID_REQUEST"):
        normalize_model_inputs({"fluid": "Oil", "unit": "m3(gas)"}, m)
