from copy import copy
from pathlib import Path
import hashlib
import openpyxl
import pytest
from openpyxl.styles import Font, PatternFill, Protection
from worker.model_preparation import PreparationError, inspect_candidate, prepare_runtime


def fixture(tmp_path):
    book=openpyxl.Workbook();sheet=book.active;sheet.title='UI'
    for row,label,value in [(1,'Primary input',5),(2,'Second input',2)]:
        sheet.cell(row,2,label);cell=sheet.cell(row,8,value)
        cell.fill=PatternFill('solid',fgColor='FFD9D9D9');cell.font=Font(color='FFC00000');cell.protection=Protection(locked=False)
    result=book.create_sheet('Report OUTPUT');result['A1']='Result';result['B1']='=UI!H1+UI!H2'
    path=tmp_path/'base.xlsx';book.save(path);book.close()
    fields=[{'id':f'demo@1.0:ui!h{row}','cell':f'H{row}','label':label,'unit':'USD','group':'inputs','type':'number','required':True,'blankOnNew':True,'default':value} for row,label,value in [(1,'Primary input',5),(2,'Second input',2)]]
    manifest={'id':'demo','version':'1.0','name':'Example','description':'Example model','inputSheet':'UI','groups':[{'id':'inputs','name':'Inputs'}], 'fields':fields,'outputs':{'views':[{'id':'results','label':'Results','sheet':'Report OUTPUT','sourceRange':'A1:B2'}]},'multipliers':[]}
    return path,manifest


def test_new_defaults_and_namespaced_ids_preserve_source(tmp_path):
    base,manifest=fixture(tmp_path);source=tmp_path/'source.xlsx'
    book=openpyxl.load_workbook(base);book['UI']['H1']=6;book.save(source);book.close()
    original=hashlib.sha256(source.read_bytes()).hexdigest()
    candidate,report=inspect_candidate(source,base,manifest,'2.0')
    assert not report['issues']
    assert candidate['fields'][0]['id']=='demo@2.0:ui!h1'
    assert candidate['fields'][0]['default']==6
    assert report['changes']==[{'label':'Primary input','kind':'default','before':5,'after':6}]
    assert hashlib.sha256(source.read_bytes()).hexdigest()==original


def test_moved_inputs_are_matched_by_fixed_labels_and_ambiguity_needs_review(tmp_path):
    base,manifest=fixture(tmp_path);source=tmp_path/'source.xlsx'
    book=openpyxl.load_workbook(base);sheet=book['UI'];style=copy(sheet['H1']._style)
    sheet['H1']=None;sheet['H1']._style=None
    sheet['I1']=5;sheet['I1']._style=style
    book.save(source);book.close()
    candidate,report=inspect_candidate(source,base,manifest,'2.0')
    assert not report['issues']
    assert candidate['fields'][0]['cell']=='I1'
    book=openpyxl.load_workbook(source);book['UI']['J1']=7;book['UI']['J1']._style=style;book.save(source);book.close()
    _,report=inspect_candidate(source,base,manifest,'2.0')
    assert report['issues']
    candidate,report=inspect_candidate(source,base,manifest,'2.0',{'fieldCells':{'demo@1.0:ui!h1':'I1'},'addedFields':[{'cell':'J1','label':'Additional amount','unit':'USD','type':'number'}]})
    assert not report['issues']
    assert len(candidate['fields'])==3
    assert candidate['fields'][-1]['group']=='additional-inputs'


def test_formula_inputs_are_blocked_and_changed_guide_requires_review(tmp_path):
    base,manifest=fixture(tmp_path);source=tmp_path/'source.xlsx'
    book=openpyxl.load_workbook(base);book.create_sheet('User Guide')['A1']='Old guidance';book.save(base)
    book['User Guide']['A1']='New guidance';book['UI']['H1']='=1+2';book.save(source);book.close()
    _,report=inspect_candidate(source,base,manifest,'2.0')
    assert report['issues'] and report['guideChanges']
    assert any('non-formula' in issue for issue in report['issues'])


def test_runtime_is_deterministic_and_only_nested_let_changes(tmp_path):
    source,_=fixture(tmp_path)
    book=openpyxl.load_workbook(source);book['Report OUTPUT']['B1']='=LET(rate,2,LET(cost,rate+1,cost*4))';book.save(source);book.close()
    original=source.read_bytes()
    first=prepare_runtime(source,tmp_path/'one.xlsx','UI')
    second=prepare_runtime(source,tmp_path/'two.xlsx','UI')
    assert first==second and first['compatibilityRewrites']==1
    assert source.read_bytes()==original
    book=openpyxl.load_workbook(tmp_path/'one.xlsx');assert book['Report OUTPUT']['B1'].value.startswith('=LET(rate,2,(');book.close()
