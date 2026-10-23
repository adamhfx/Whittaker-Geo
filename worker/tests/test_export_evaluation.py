import base64
import json
import uuid
import zipfile

import openpyxl
import pytest

from worker.export_evaluation import export_evaluation
from worker.live import PreviewService, LiveError
from worker.model_preview import PreviewError


def fixture(tmp_path):
    template=tmp_path/'source.xlsx'
    book=openpyxl.Workbook()
    book.active.title='Private MODEL'
    book.active['A1']='PRIVATE_VALUE_MUST_NOT_LEAK'
    output=book.create_sheet('Summary OUTPUT')
    output['A1']='OLD_CACHED_OUTPUT'
    output['A2']='HIDDEN_ROW_MUST_NOT_LEAK'
    output['B1']='HIDDEN_COLUMN_MUST_NOT_LEAK'
    output.row_dimensions[2].hidden=True
    output.column_dimensions['B'].hidden=True
    book.save(template); book.close()
    manifest={'id':'test','version':'v2','outputs':{'views':[
        {'id':'summary','label':'Summary','sheet':'Summary OUTPUT','sourceRange':'A1:C3'}]}}
    result={'calculatorId':'test','calculatorVersion':'v2','durationSeconds':0.1,'outputSheets':[
        {'name':'Summary OUTPUT','rows':[
            {'cells':[{'column':1,'value':'=LITERAL_TEXT'},{'column':3,'value':125.5,'format':'$0.00'}]},
            {'cells':[{'column':1,'value':0},{'column':3,'value':None}]}]}]}
    return template,manifest,result


def test_export_uses_current_values_without_private_cells_or_formulas(tmp_path):
    template,manifest,result=fixture(tmp_path)
    target=tmp_path/'export.xlsx'
    export_evaluation(template,manifest,result,target)
    book=openpyxl.load_workbook(target,data_only=False)
    assert book.sheetnames==['Summary OUTPUT']
    sheet=book.active
    assert sheet['A1'].value=='=LITERAL_TEXT' and sheet['A1'].data_type=='s'
    assert sheet['B1'].value==125.5 and sheet['B1'].number_format=='$0.00'
    assert sheet['A2'].value==0
    assert sheet['C1'].value is None and sheet['A3'].value is None
    assert sheet['A1'].border.left.style=='thin' and sheet['B2'].border.bottom.style=='thin'
    book.close()
    with zipfile.ZipFile(target) as archive:
        data=b''.join(archive.read(name) for name in archive.namelist())
        for forbidden in (b'PRIVATE_VALUE',b'HIDDEN_ROW',b'HIDDEN_COLUMN',b'OLD_CACHED_OUTPUT'):
            assert forbidden not in data


def test_export_rejects_mismatched_output_geometry(tmp_path):
    template,manifest,result=fixture(tmp_path)
    result['outputSheets'][0]['rows'][0]['cells'][1]['column']=2
    with pytest.raises(PreviewError,match='RESULT_INVALID'):
        export_evaluation(template,manifest,result,tmp_path/'export.xlsx')


def test_export_endpoint_shares_capacity_and_returns_validated_values(tmp_path):
    template,manifest,result=fixture(tmp_path)
    def loader(calculator,version,target,**kwargs):
        target.write_bytes(template.read_bytes())
        return manifest
    def evaluate(*args,**kwargs):
        assert kwargs['timeout']==90
        return result
    service=PreviewService(engine_secret='a'*32,supabase_url='https://example.test',service_key='server-only',model_loader=loader,evaluator=evaluate)
    raw=json.dumps({'calculatorId':'test','version':'v2','inputs':{},'requesterId':str(uuid.uuid4())}).encode()
    payload=service.evaluate(raw,export=True)
    assert payload['result']==result
    assert base64.b64decode(payload['xlsxBase64']).startswith(b'PK\x03\x04')
    service._capacity.acquire()
    try:
        with pytest.raises(LiveError) as caught: service.evaluate(raw,export=True)
        assert caught.value.code=='ENGINE_BUSY'
    finally: service._capacity.release()


def test_offset_output_starts_at_a1_with_borders_and_readable_headers(tmp_path):
    template=tmp_path/'offset.xlsx'
    book=openpyxl.Workbook();book.active.title='Annual OUTPUT'
    book.active['N9']='Annual results';book.active['N11']='Year';book.active['O11']='Gross revenue';book.active['N12']=1;book.active['O12']=123.45
    book.save(template);book.close()
    manifest={'id':'test','version':'v2','outputs':{'views':[{'id':'annual','label':'Annual','sheet':'Annual OUTPUT','sourceRange':'N9:P12'}]}}
    result={'outputSheets':[{'name':'Annual OUTPUT','rows':[
        {'cells':[{'column':14,'value':'Annual results'}]}, {'cells':[]},
        {'cells':[{'column':14,'value':'Year'},{'column':15,'value':'Gross revenue'}]},
        {'cells':[{'column':14,'value':1},{'column':15,'value':123.45,'format':'$0.00'}]}]}]}
    target=tmp_path/'offset-export.xlsx';export_evaluation(template,manifest,result,target)
    exported=openpyxl.load_workbook(target);sheet=exported.active
    assert sheet['A1'].value=='Annual results' and sheet['B4'].value==123.45
    assert sheet['B4'].number_format=='$0.00'
    assert sheet['A4'].border.bottom.style=='thin'
    assert sheet['A3'].font.bold and sheet['A3'].alignment.wrap_text
    assert sheet.freeze_panes=='B4'
    assert sheet['N9'].value is None
    exported.close()
